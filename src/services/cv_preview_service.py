from __future__ import annotations

import base64
import html
from pathlib import Path

import pymupdf
from docx import Document

from services.settings_service import SettingsService


class CVPreviewService:
    """Create a local, browser-friendly preview of the CV selected in JAM."""

    def __init__(self) -> None:
        self._settings = SettingsService()

    def preview(self) -> dict:
        path_value = self._settings.get("cv_path", "").strip()

        if not path_value:
            raise ValueError("Choose a CV first.")

        path = Path(path_value)

        if not path.exists():
            raise FileNotFoundError("The selected CV file no longer exists.")

        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return self._preview_pdf(path)

        if suffix == ".docx":
            return self._preview_docx(path)

        if suffix in {".txt", ".md", ".cvm"}:
            return self._preview_text(path)

        raise ValueError("Preview supports PDF, DOCX, TXT, MD and CVM files.")

    @staticmethod
    def _preview_pdf(path: Path) -> dict:
        pages: list[str] = []

        with pymupdf.open(path) as document:
            page_count = min(document.page_count, 12)
            matrix = pymupdf.Matrix(1.55, 1.55)

            for index in range(page_count):
                page = document.load_page(index)
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                encoded = base64.b64encode(pix.tobytes("png")).decode("ascii")
                pages.append(f"data:image/png;base64,{encoded}")

            truncated = document.page_count > page_count

        return {
            "kind": "pdf",
            "name": path.name,
            "path": str(path),
            "pages": pages,
            "page_count": len(pages),
            "truncated": truncated,
        }

    @staticmethod
    def _preview_docx(path: Path) -> dict:
        document = Document(path)
        chunks: list[str] = []

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if value:
                chunks.append(value)

        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    chunks.append(" | ".join(values))

        return {
            "kind": "text",
            "name": path.name,
            "path": str(path),
            "text": "\n\n".join(chunks),
            "page_count": None,
            "truncated": False,
        }

    @staticmethod
    def _preview_text(path: Path) -> dict:
        raw = path.read_text(encoding="utf-8", errors="ignore")

        if path.suffix.lower() == ".cvm":
            try:
                import json

                # Import this helper only for CVM previews. Keeping the import
                # out of module startup prevents CV preview support from
                # pulling the full Analyzer stack into JAM at launch.
                from services.analyzer_service import flatten_json_strings

                raw = flatten_json_strings(json.loads(raw))
            except Exception:
                pass

        return {
            "kind": "text",
            "name": path.name,
            "path": str(path),
            "text": raw,
            "page_count": None,
            "truncated": False,
        }
