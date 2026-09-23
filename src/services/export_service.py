from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from PIL import Image as PILImage
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from config import BRANDING_DIR, EXPORT_DIR
from database import connection, utc_now
from services.settings_service import SettingsService


GOLD = "BC965D"
GOLD_LIGHT = "D0AA70"
NAVY = "111827"
NAVY_DARK = "0B111A"
GRAPHITE = "374151"
MUTED = "6B7280"
SOFT_ICE = "F7F9FC"
SOFT_GOLD = "F6EFE3"
WHITE = "FFFFFF"
LINE = "D9DEE7"

STATUS_COLORS = {
    "Saved": "E5E7EB",
    "To Review": "E9D5FF",
    "Applied": "FEF3C7",
    "Interviewing": "DBEAFE",
    "Offer": "FFEDD5",
    "Accepted": "D1FAE5",
    "Rejected": "FEE2E2",
    "Withdrawn": "E5E7EB",
    "Archived": "E5E7EB",
}


EXPORT_FR = {
    "APPLICATIONS EXPORT": "EXPORT DES CANDIDATURES",
    "FILTERS APPLIED": "FILTRES APPLIQUÉS",
    "Date Saved": "Date d’enregistrement",
    "Company": "Entreprise",
    "Job Title": "Poste",
    "Status": "Statut",
    "Match Score": "Score de correspondance",
    "Rating": "Note",
    "Location": "Localisation",
    "Work Mode": "Mode de travail",
    "Source": "Source",
    "Date Applied": "Date de candidature",
    "Contact": "Contact",
    "Contact Link": "Lien du contact",
    "Job URL": "URL de l’offre",
    "Notes": "Notes",
    "JAM — Export Summary": "JAM — Résumé de l’export",
    "TOTAL APPLICATIONS": "TOTAL DES CANDIDATURES",
    "INTERVIEWING": "ENTRETIENS",
    "ACCEPTED": "ACCEPTÉES",
    "AVERAGE MATCH SCORE": "SCORE MOYEN",
    "FILTERS": "FILTRES",
    "EXPORTED ROWS": "LIGNES EXPORTÉES",
    "Applications": "Candidatures",
    "Summary": "Résumé",
    "Saved": "Enregistrée",
    "To Review": "À examiner",
    "Applied": "Candidature envoyée",
    "Interviewing": "Entretien",
    "Offer": "Offre reçue",
    "Accepted": "Acceptée",
    "Rejected": "Refusée",
    "Withdrawn": "Retirée",
    "Archived": "Archivée",
    "Onsite": "Sur site",
    "Hybrid": "Hybride",
    "Remote": "Télétravail",
    "Filters applied": "Filtres appliqués",
    "Score": "Score",
    "No filters applied — all visible applications were exported.":
        "Aucun filtre appliqué — toutes les candidatures visibles ont été exportées.",
}


class ExportService:
    def __init__(self) -> None:
        self._settings = SettingsService()

    def _language(self) -> str:
        settings = getattr(self, "_settings", None)
        if settings is None:
            return "en"
        language = settings.get("language", "en").strip().lower()
        return language if language in {"en", "fr"} else "en"

    def _t(self, text: str) -> str:
        if self._language() == "fr":
            return EXPORT_FR.get(text, text)
        return text

    def _status_label(self, status: str) -> str:
        return self._t(status)

    def _work_mode_label(self, value: str) -> str:
        return self._t(str(value or ""))

    def _human_date(self, value: datetime) -> str:
        if self._language() != "fr":
            return value.strftime("%d %b %Y")

        months = (
            "janv.", "févr.", "mars", "avr.", "mai", "juin",
            "juil.", "août", "sept.", "oct.", "nov.", "déc.",
        )
        return f"{value.day:02d} {months[value.month - 1]} {value.year}"

    def _human_datetime(self, value: datetime) -> str:
        return f"{self._human_date(value)} • {value.strftime('%H:%M')}"

    def _translate_filter_line(self, text: str) -> str:
        if self._language() != "fr":
            return text

        value = str(text)

        exact = {
            "No filters applied — all visible applications.":
                "Aucun filtre appliqué — toutes les candidatures visibles.",
            "No filters applied — all visible applications were exported.":
                "Aucun filtre appliqué — toutes les candidatures visibles ont été exportées.",
        }

        if value in exact:
            return exact[value]

        prefix_map = {
            "Search:": "Recherche :",
            "Company:": "Entreprise :",
            "Job title:": "Poste :",
            "Status:": "Statut :",
            "Score:": "Score :",
            "Rating:": "Note :",
            "Location:": "Localisation :",
            "Date saved:": "Date d’enregistrement :",
            "Source:": "Source :",
            "Sort:": "Tri :",
            "Manual selection:": "Sélection manuelle :",
        }

        for prefix, translated in prefix_map.items():
            if value.startswith(prefix):
                value = translated + value[len(prefix):]
                break

        replacements = {
            "ascending": "croissant",
            "descending": "décroissant",
            "None": "Aucun",
            "Saved": "Enregistrée",
            "To Review": "À examiner",
            "Applied": "Candidature envoyée",
            "Interviewing": "Entretien",
            "Offer": "Offre reçue",
            "Accepted": "Acceptée",
            "Rejected": "Refusée",
            "Withdrawn": "Retirée",
            "Archived": "Archivée",
            "application(s)": "candidature(s)",
            "No rating": "Aucune note",
            "No score": "Aucun score",
        }

        for source, target in replacements.items():
            value = value.replace(source, target)

        return value

    def _filename(self, suffix: str) -> Path:
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        return EXPORT_DIR / f"JAM_Applications_{stamp}.{suffix}"

    def _track(self, export_type: str, path: Path, row_count: int) -> None:
        with connection() as conn:
            conn.execute(
                "INSERT INTO export_history(export_type, path, row_count, created_at) VALUES (?, ?, ?, ?)",
                (export_type, str(path), row_count, utc_now()),
            )

    def history(self) -> list[dict]:
        with connection() as conn:
            rows = conn.execute(
                "SELECT * FROM export_history ORDER BY datetime(created_at) DESC, id DESC"
            ).fetchall()

        return [dict(row) for row in rows]

    def remove_history(self, history_id: int) -> bool:
        with connection() as conn:
            cursor = conn.execute(
                "DELETE FROM export_history WHERE id = ?",
                (int(history_id),),
            )
            return cursor.rowcount > 0

    def _filters_text(self, filter_summary: list[str] | None) -> list[str]:
        clean = [
            self._translate_filter_line(str(item).strip())
            for item in (filter_summary or [])
            if str(item).strip()
        ]

        return clean or [
            self._t(
                "No filters applied — all visible applications were exported."
            )
        ]

    @staticmethod
    def _excel_datetime(value):
        """
        Convert JAM ISO timestamps into real Excel datetimes.

        Excel cannot store timezone-aware datetimes directly, so the displayed
        local clock value is preserved and timezone info is removed.
        """
        if value in (None, ""):
            return None

        if isinstance(value, datetime):
            return value.replace(tzinfo=None)

        text = str(value).strip()

        if not text:
            return None

        try:
            parsed = datetime.fromisoformat(
                text.replace("Z", "+00:00")
            )
            return parsed.replace(tzinfo=None)
        except Exception:
            return text

    @staticmethod
    def _center_merged(ws, cell_range: str, value, *, font, fill=None, row_height=None):
        ws.merge_cells(cell_range)

        start = cell_range.split(":")[0]
        cell = ws[start]
        cell.value = value
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        if fill:
            fill_obj = PatternFill("solid", fgColor=fill)
            for row in ws[cell_range]:
                for merged_cell in row:
                    merged_cell.fill = fill_obj

        if row_height:
            first_row = ws[start].row
            last_row = ws[cell_range.split(":")[1]].row

            for row_number in range(first_row, last_row + 1):
                ws.row_dimensions[row_number].height = row_height

    def _add_excel_branding(
        self,
        ws,
        row_count: int,
        filter_summary: list[str] | None,
    ) -> int:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.outlinePr.summaryBelow = True

        # ------------------------------------------------------------
        # HERO HEADER
        # ------------------------------------------------------------
        ws.merge_cells("A1:N2")

        hero = ws["A1"]
        hero.value = "JAM — Job Application Manager"
        hero.font = Font(
            name="Aptos Display",
            size=24,
            bold=True,
            color=WHITE,
        )
        hero.fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )
        hero.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        # Fill the complete merged banner.
        for row in ws["A1:N2"]:
            for cell in row:
                cell.fill = PatternFill(
                    "solid",
                    fgColor=NAVY,
                )

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 28

        logo_path = BRANDING_DIR / "jam_logo.png"

        if logo_path.exists():
            try:
                logo = XLImage(str(logo_path))
                logo.width = 70
                logo.height = 44
                ws.add_image(
                    logo,
                    "A1",
                )
            except Exception:
                pass

        generated_at = datetime.now()

        self._center_merged(
            ws,
            "A3:N3",
            self._t("APPLICATIONS EXPORT"),
            font=Font(
                name="Aptos",
                size=11,
                bold=True,
                color=GOLD_LIGHT,
            ),
            fill=NAVY,
            row_height=22,
        )

        self._center_merged(
            ws,
            "A4:N4",
            (
                (
                    f"{row_count} candidature"
                    f"{'' if row_count == 1 else 's'}"
                    f"  •  Généré le {self._human_date(generated_at)} à {generated_at.strftime('%H:%M')}"
                    if self._language() == "fr"
                    else
                    f"{row_count} application"
                    f"{'' if row_count == 1 else 's'}"
                    f"  •  Generated {generated_at.strftime('%d %b %Y')} at "
                    f"{generated_at.strftime('%H:%M')}"
                )
            ),
            font=Font(
                name="Aptos",
                size=10,
                color="D1D5DB",
            ),
            fill=NAVY,
            row_height=20,
        )

        ws.row_dimensions[5].height = 9

        # ------------------------------------------------------------
        # FILTER SUMMARY
        # ------------------------------------------------------------
        self._center_merged(
            ws,
            "A6:N6",
            self._t("FILTERS APPLIED"),
            font=Font(
                name="Aptos",
                size=10,
                bold=True,
                color=GOLD,
            ),
            fill=SOFT_GOLD,
            row_height=22,
        )

        filters = self._filters_text(
            filter_summary
        )

        row = 7

        for item in filters[:8]:
            self._center_merged(
                ws,
                f"A{row}:N{row}",
                f"• {item}",
                font=Font(
                    name="Aptos",
                    size=9,
                    color=GRAPHITE,
                ),
                fill=WHITE,
                row_height=20,
            )
            row += 1

        ws.row_dimensions[row].height = 9

        return row + 1

    def export_excel(
        self,
        rows: list[dict],
        filter_summary: list[str] | None = None,
    ) -> str:
        path = self._filename("xlsx")

        workbook = Workbook()

        ws = workbook.active
        ws.title = self._t("Applications")

        header_row = self._add_excel_branding(
            ws,
            len(rows),
            filter_summary,
        )

        headers = [
            self._t("Date Saved"),
            self._t("Company"),
            self._t("Job Title"),
            self._t("Status"),
            self._t("Match Score"),
            self._t("Rating"),
            self._t("Location"),
            self._t("Work Mode"),
            self._t("Source"),
            self._t("Date Applied"),
            self._t("Contact"),
            self._t("Contact Link"),
            self._t("Job URL"),
            self._t("Notes"),
        ]

        for col, value in enumerate(
            headers,
            start=1,
        ):
            cell = ws.cell(
                row=header_row,
                column=col,
                value=value,
            )
            cell.fill = PatternFill(
                "solid",
                fgColor=NAVY,
            )
            cell.font = Font(
                name="Aptos",
                color=WHITE,
                bold=True,
                size=10,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(
                bottom=Side(
                    style="medium",
                    color=GOLD,
                )
            )

        ws.row_dimensions[
            header_row
        ].height = 30

        data_start = (
            header_row + 1
        )

        thin_line = Side(
            style="hair",
            color=LINE,
        )

        for row_idx, row in enumerate(
            rows,
            start=data_start,
        ):
            date_saved = self._excel_datetime(
                row.get("date_saved")
            )
            date_applied = self._excel_datetime(
                row.get("date_applied")
            )

            score = (
                None
                if row.get("match_score") is None
                else int(
                    row.get("match_score")
                )
            )

            rating = int(
                row.get("rating") or 0
            )

            values = [
                date_saved,
                row.get("company", ""),
                row.get("job_title", ""),
                self._status_label(row.get("status", "")),
                score,
                "★" * rating if rating else "",
                row.get("location", ""),
                self._work_mode_label(row.get("work_mode", "")),
                row.get("source", ""),
                date_applied,
                row.get("contact_name", ""),
                row.get("contact_url", ""),
                row.get("url", ""),
                row.get("notes", ""),
            ]

            base_fill = (
                WHITE
                if (
                    row_idx - data_start
                ) % 2 == 0
                else SOFT_ICE
            )

            for col, value in enumerate(
                values,
                start=1,
            ):
                cell = ws.cell(
                    row=row_idx,
                    column=col,
                    value=value,
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor=base_fill,
                )

                cell.font = Font(
                    name="Aptos",
                    size=9,
                    color=NAVY,
                )

                cell.border = Border(
                    bottom=thin_line,
                )

                # Center compact categorical fields and dates.
                if col in {
                    1,
                    4,
                    5,
                    6,
                    7,
                    8,
                    9,
                    10,
                }:
                    horizontal = "center"
                else:
                    horizontal = "left"

                cell.alignment = Alignment(
                    horizontal=horizontal,
                    vertical="center",
                    wrap_text=True,
                )

            # Real Excel date formatting.
            for date_col in (
                1,
                10,
            ):
                date_cell = ws.cell(
                    row=row_idx,
                    column=date_col,
                )

                if isinstance(
                    date_cell.value,
                    datetime,
                ):
                    date_cell.number_format = (
                        "[$-fr-FR]dd mmm yyyy  hh:mm"
                        if self._language() == "fr"
                        else "dd mmm yyyy  hh:mm"
                    )

            # Status chip-like visual.
            status_cell = ws.cell(
                row=row_idx,
                column=4,
            )
            status_cell.font = Font(
                name="Aptos",
                bold=True,
                size=9,
                color=NAVY,
            )
            status_cell.fill = PatternFill(
                "solid",
                fgColor=STATUS_COLORS.get(
                    str(
                        row.get("status")
                        or ""
                    ),
                    base_fill,
                ),
            )

            # Score as a real numeric percentage-looking value.
            score_cell = ws.cell(
                row=row_idx,
                column=5,
            )

            if score is not None:
                score_cell.number_format = (
                    '0"%"'
                )
                score_cell.font = Font(
                    name="Aptos",
                    bold=True,
                    size=9,
                    color=GOLD,
                )

            rating_cell = ws.cell(
                row=row_idx,
                column=6,
            )
            rating_cell.font = Font(
                name="Segoe UI Symbol",
                bold=True,
                size=10,
                color=GOLD,
            )

            # Clickable URLs.
            for url_col in (
                12,
                13,
            ):
                url_cell = ws.cell(
                    row=row_idx,
                    column=url_col,
                )

                if str(
                    url_cell.value
                    or ""
                ).startswith(
                    (
                        "http://",
                        "https://",
                    )
                ):
                    url_cell.hyperlink = (
                        str(
                            url_cell.value
                        )
                    )
                    url_cell.style = (
                        "Hyperlink"
                    )
                    url_cell.alignment = Alignment(
                        horizontal="left",
                        vertical="center",
                        wrap_text=True,
                    )

            ws.row_dimensions[
                row_idx
            ].height = 32

        widths = [
            21,
            24,
            34,
            17,
            14,
            14,
            22,
            16,
            24,
            21,
            22,
            31,
            38,
            42,
        ]

        for index, width in enumerate(
            widths,
            start=1,
        ):
            ws.column_dimensions[
                get_column_letter(index)
            ].width = width

        last_row = max(
            header_row,
            data_start + len(rows) - 1,
        )

        # Standard worksheet filter only.
        # No Structured Table is created, avoiding Excel repair warnings.
        ws.auto_filter.ref = (
            f"A{header_row}:N{last_row}"
        )

        ws.freeze_panes = (
            f"A{data_start}"
        )

        # Gold score data bars while keeping the numeric score visible.
        if len(rows):
            ws.conditional_formatting.add(
                f"E{data_start}:E{last_row}",
                DataBarRule(
                    start_type="num",
                    start_value=0,
                    end_type="num",
                    end_value=100,
                    color=GOLD,
                    showValue=True,
                ),
            )

        # Print/export polish.
        ws.print_title_rows = (
            f"{header_row}:{header_row}"
        )
        ws.print_area = (
            f"A1:N{last_row}"
        )
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4

        # ------------------------------------------------------------
        # SUMMARY SHEET
        # ------------------------------------------------------------
        summary = workbook.create_sheet(
            self._t("Summary")
        )
        summary.sheet_view.showGridLines = False

        summary.merge_cells(
            "A1:F2"
        )
        summary["A1"] = (
            self._t("JAM — Export Summary")
        )
        summary["A1"].font = Font(
            name="Aptos Display",
            size=22,
            bold=True,
            color=WHITE,
        )
        summary["A1"].fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )
        summary["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        for row in summary[
            "A1:F2"
        ]:
            for cell in row:
                cell.fill = PatternFill(
                    "solid",
                    fgColor=NAVY,
                )

        summary.row_dimensions[1].height = 28
        summary.row_dimensions[2].height = 24

        self._center_merged(
            summary,
            "A3:F3",
            (
                (
                    f"Généré le {self._human_datetime(datetime.now())}"
                    if self._language() == "fr"
                    else
                    f"Generated {datetime.now().strftime('%d %b %Y • %H:%M')}"
                )
            ),
            font=Font(
                name="Aptos",
                size=10,
                color=GOLD_LIGHT,
            ),
            fill=NAVY,
            row_height=20,
        )

        statuses: dict[str, int] = {}

        scores = []

        for row in rows:
            status = (
                row.get("status")
                or "Unknown"
            )

            statuses[status] = (
                statuses.get(
                    status,
                    0,
                ) + 1
            )

            if (
                row.get(
                    "match_score"
                )
                is not None
            ):
                scores.append(
                    int(
                        row[
                            "match_score"
                        ]
                    )
                )

        accepted_count = (
            statuses.get(
                "Accepted",
                0,
            )
        )

        interview_count = (
            statuses.get(
                "Interviewing",
                0,
            )
        )

        summary.merge_cells(
            "A5:B5"
        )
        summary.merge_cells(
            "C5:D5"
        )
        summary.merge_cells(
            "E5:F5"
        )

        summary[
            "A5"
        ] = (
            f"{self._t('TOTAL APPLICATIONS')}\n"
            f"{len(rows)}"
        )

        summary[
            "C5"
        ] = (
            f"{self._t('INTERVIEWING')}\n"
            f"{interview_count}"
        )

        summary[
            "E5"
        ] = (
            f"{self._t('ACCEPTED')}\n"
            f"{accepted_count}"
        )

        for coordinate in (
            "A5",
            "C5",
            "E5",
        ):
            cell = summary[
                coordinate
            ]
            cell.fill = PatternFill(
                "solid",
                fgColor=SOFT_GOLD,
            )
            cell.font = Font(
                name="Aptos",
                size=11,
                bold=True,
                color=NAVY,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        summary.row_dimensions[
            5
        ].height = 52

        summary.merge_cells(
            "A7:B7"
        )
        summary.merge_cells(
            "C7:D7"
        )
        summary.merge_cells(
            "E7:F7"
        )

        summary[
            "A7"
        ] = (
            f"{self._t('AVERAGE MATCH SCORE')}\n"
            + (
                f"{round(sum(scores) / len(scores), 1)}%"
                if scores
                else "—"
            )
        )

        summary[
            "C7"
        ] = (
            f"{self._t('FILTERS')}\n"
            f"{len(self._filters_text(filter_summary))}"
        )

        summary[
            "E7"
        ] = (
            f"{self._t('EXPORTED ROWS')}\n"
            f"{len(rows)}"
        )

        for coordinate in (
            "A7",
            "C7",
            "E7",
        ):
            cell = summary[
                coordinate
            ]
            cell.fill = PatternFill(
                "solid",
                fgColor=NAVY,
            )
            cell.font = Font(
                name="Aptos",
                size=10,
                bold=True,
                color=WHITE,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        summary.row_dimensions[
            7
        ].height = 48

        self._center_merged(
            summary,
            "A9:F9",
            self._t("FILTERS APPLIED"),
            font=Font(
                name="Aptos",
                size=10,
                bold=True,
                color=GOLD,
            ),
            fill=SOFT_GOLD,
            row_height=22,
        )

        summary_row = 10

        for item in self._filters_text(
            filter_summary
        ):
            self._center_merged(
                summary,
                f"A{summary_row}:F{summary_row}",
                f"• {item}",
                font=Font(
                    name="Aptos",
                    size=9,
                    color=GRAPHITE,
                ),
                fill=WHITE,
                row_height=20,
            )
            summary_row += 1

        summary_row += 1

        summary.cell(
            row=summary_row,
            column=1,
            value=self._t("Status"),
        )

        summary.cell(
            row=summary_row,
            column=2,
            value=self._t("Applications"),
        )

        for cell in summary[
            summary_row
        ][:2]:
            cell.fill = PatternFill(
                "solid",
                fgColor=NAVY,
            )
            cell.font = Font(
                name="Aptos",
                color=WHITE,
                bold=True,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for status, count in sorted(
            statuses.items()
        ):
            summary_row += 1

            summary.cell(
                row=summary_row,
                column=1,
                value=self._status_label(status),
            )

            summary.cell(
                row=summary_row,
                column=2,
                value=count,
            )

            summary.cell(
                row=summary_row,
                column=1,
            ).fill = PatternFill(
                "solid",
                fgColor=STATUS_COLORS.get(
                    status,
                    WHITE,
                ),
            )

            summary.cell(
                row=summary_row,
                column=1,
            ).font = Font(
                name="Aptos",
                bold=True,
                color=NAVY,
            )

            summary.cell(
                row=summary_row,
                column=2,
            ).alignment = Alignment(
                horizontal="center",
            )

        for column in (
            "A",
            "B",
            "C",
            "D",
            "E",
            "F",
        ):
            summary.column_dimensions[
                column
            ].width = 22

        workbook.save(
            path
        )

        self._track(
            "Excel",
            path,
            len(rows),
        )

        return str(path)

    def export_pdf(
        self,
        rows: list[dict],
        filter_summary: list[str] | None = None,
    ) -> str:
        path = self._filename("pdf")

        doc = SimpleDocTemplate(
            str(path),
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=11 * mm,
            bottomMargin=11 * mm,
            title=(
                "JAM — Export des candidatures"
                if self._language() == "fr"
                else "JAM Applications Export"
            ),
            author="JAM - By Farouk",
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "JAMTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=19,
            textColor=colors.HexColor(
                f"#{NAVY}"
            ),
            alignment=TA_CENTER,
            spaceAfter=2 * mm,
        )

        meta_style = ParagraphStyle(
            "JAMMeta",
            parent=styles["BodyText"],
            fontSize=8.5,
            textColor=colors.HexColor(
                f"#{GRAPHITE}"
            ),
            leading=11,
            alignment=TA_CENTER,
        )

        filter_style = ParagraphStyle(
            "JAMFilter",
            parent=meta_style,
            fontSize=7.6,
            leading=9.5,
            textColor=colors.HexColor(
                f"#{GRAPHITE}"
            ),
            alignment=TA_CENTER,
        )

        cell_style = ParagraphStyle(
            "JAMCell",
            parent=styles["BodyText"],
            fontSize=7.2,
            leading=9,
            textColor=colors.HexColor(
                f"#{NAVY}"
            ),
        )

        head_style = ParagraphStyle(
            "JAMHead",
            parent=cell_style,
            fontName="Helvetica-Bold",
            textColor=colors.white,
            alignment=TA_CENTER,
        )

        story = []

        logo_path = BRANDING_DIR / "jam_logo.png"

        if logo_path.exists():
            try:
                with PILImage.open(logo_path) as image:
                    pixel_width, pixel_height = image.size

                target_width = 100 * mm
                aspect = (pixel_height / pixel_width) if pixel_width else 0.5
                target_height = target_width * aspect

                # Keep the real logo proportions and give the brand proper
                # visual weight without letting an unusually tall asset take
                # over the first page.
                max_height = 50 * mm
                if target_height > max_height:
                    target_height = max_height
                    target_width = target_height / max(aspect, 0.01)

                logo = RLImage(
                    str(logo_path),
                    width=target_width,
                    height=target_height,
                )
                logo.hAlign = "CENTER"
                story.append(logo)
                story.append(
                    Spacer(
                        1,
                        3 * mm,
                    )
                )
            except Exception:
                pass

        story.extend(
            [
                Paragraph(
                    "JAM — Job Application Manager",
                    title_style,
                ),
                Paragraph(
                    (
                        (
                            f"Export des candidatures • "
                            f"{self._human_datetime(datetime.now())} • "
                            f"{len(rows)} candidature"
                            f"{'' if len(rows) == 1 else 's'}"
                            if self._language() == "fr"
                            else
                            f"Applications export • "
                            f"{datetime.now().strftime('%d %b %Y • %H:%M')} • "
                            f"{len(rows)} application"
                            f"{'' if len(rows) == 1 else 's'}"
                        )
                    ),
                    meta_style,
                ),
                Spacer(
                    1,
                    2 * mm,
                ),
                Paragraph(
                    f"<b>{self._t('Filters applied')}</b>",
                    meta_style,
                ),
            ]
        )

        for item in self._filters_text(
            filter_summary
        ):
            story.append(
                Paragraph(
                    f"• {item}",
                    filter_style,
                )
            )

        story.append(
            Spacer(
                1,
                4 * mm,
            )
        )

        headers = [
            self._t("Date Saved"),
            self._t("Company"),
            self._t("Job Title"),
            self._t("Status"),
            self._t("Score"),
            self._t("Rating"),
            self._t("Location"),
            self._t("Source"),
            self._t("Applied"),
        ]

        data = [
            [
                Paragraph(
                    value,
                    head_style,
                )
                for value in headers
            ]
        ]

        for row in rows:
            rating = int(
                row.get("rating")
                or 0
            )

            saved = self._excel_datetime(
                row.get("date_saved")
            )

            applied = self._excel_datetime(
                row.get("date_applied")
            )

            saved_text = (
                self._human_date(saved)
                if isinstance(
                    saved,
                    datetime,
                )
                else (
                    str(saved)
                    if saved
                    else "—"
                )
            )

            applied_text = (
                self._human_date(applied)
                if isinstance(
                    applied,
                    datetime,
                )
                else (
                    str(applied)
                    if applied
                    else "—"
                )
            )

            values = [
                saved_text,
                row.get(
                    "company",
                    "",
                ),
                row.get(
                    "job_title",
                    "",
                ),
                self._status_label(
                    row.get(
                        "status",
                        "",
                    )
                ),
                (
                    "—"
                    if row.get(
                        "match_score"
                    )
                    is None
                    else f"{row.get('match_score')}%"
                ),
                (
                    f"{rating}/5"
                    if rating
                    else "—"
                ),
                row.get(
                    "location",
                    "",
                ),
                row.get(
                    "source",
                    "",
                ),
                applied_text,
            ]

            data.append(
                [
                    Paragraph(
                        str(value),
                        cell_style,
                    )
                    for value in values
                ]
            )

        table = Table(
            data,
            repeatRows=1,
            colWidths=[
                20 * mm,
                34 * mm,
                52 * mm,
                27 * mm,
                18 * mm,
                22 * mm,
                34 * mm,
                27 * mm,
                20 * mm,
            ],
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            f"#{NAVY}"
                        ),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, 0),
                        "CENTER",
                    ),
                    (
                        "LINEBELOW",
                        (0, 0),
                        (-1, 0),
                        1.2,
                        colors.HexColor(
                            f"#{GOLD}"
                        ),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.25,
                        colors.HexColor(
                            "#D7DCE4"
                        ),
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor(
                                f"#{SOFT_ICE}"
                            ),
                        ],
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(
            table
        )

        doc.build(
            story
        )

        self._track(
            "PDF",
            path,
            len(rows),
        )

        return str(path)
