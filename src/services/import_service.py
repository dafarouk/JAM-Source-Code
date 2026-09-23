from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from repositories.application_repository import ApplicationRepository


HEADER_ALIASES = {
    "company": {"company", "company name", "entreprise", "societe", "société", "employer", "employeur"},
    "job_title": {"job title", "title", "role", "position", "poste", "intitule", "intitulé"},
    "status": {"status", "statut", "application status"},
    "match_score": {"match score", "score", "jam score", "cv match", "matching score"},
    "rating": {"rating", "stars", "star rating", "note", "priority"},
    "location": {"location", "city", "ville", "lieu", "localisation"},
    "work_mode": {"work mode", "mode", "remote", "work type", "working mode", "mode de travail"},
    "source": {"source", "job board", "platform", "site", "website"},
    "date_saved": {"date saved", "saved date", "date", "date added", "created", "created at"},
    "date_applied": {"date applied", "application date", "applied date", "date candidature"},
    "contact_name": {"contact", "contact name", "recruiter", "recruiter name", "rh", "hr contact"},
    "contact_url": {"contact link", "contact url", "recruiter link", "linkedin contact"},
    "url": {"job url", "url", "link", "job link", "offer link", "lien"},
    "salary_text": {"salary", "salary text", "compensation", "salaire", "pay"},
    "salary_min": {"salary min", "minimum salary", "min salary", "salaire min", "salaire minimum"},
    "salary_max": {"salary max", "maximum salary", "max salary", "salaire max", "salaire maximum"},
    "salary_currency": {"salary currency", "currency", "devise", "devise salaire"},
    "salary_period": {"salary period", "period", "periode", "période", "monthly annual"},
    "salary_basis": {"salary basis", "gross net", "brut net", "basis"},
    "description": {"description", "job description", "offer description", "jd"},
    "notes": {"notes", "note", "comments", "comment", "remarks", "remark"},
}


def _norm(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


NORMALIZED_ALIASES = {
    field: {_norm(alias) for alias in aliases}
    for field, aliases in HEADER_ALIASES.items()
}


class ExcelImportService:
    def __init__(self, repository: ApplicationRepository | None = None) -> None:
        self.repo = repository or ApplicationRepository()

    @staticmethod
    def _value_to_text(value) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.astimezone().isoformat(timespec="seconds") if value.tzinfo else value.isoformat(timespec="seconds")
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time()).isoformat(timespec="seconds")
        return str(value).strip()

    @staticmethod
    def _rating(value) -> int:
        if value is None:
            return 0
        text = str(value).strip()
        stars = text.count("★")
        if stars:
            return max(0, min(5, stars))
        match = re.search(r"\d+(?:[.,]\d+)?", text)
        if not match:
            return 0
        return max(0, min(5, int(float(match.group(0).replace(",", ".")))))

    @staticmethod
    def _score(value):
        if value is None or str(value).strip() == "":
            return None
        match = re.search(r"-?\d+(?:[.,]\d+)?", str(value))
        if not match:
            return None
        return max(0, min(100, int(round(float(match.group(0).replace(",", "."))))))

    @staticmethod
    def _number(value):
        if value is None or str(value).strip() == "":
            return None
        match = re.search(r"-?\d+(?:[.,]\d+)?", str(value).replace(" ", ""))
        if not match:
            return None
        return float(match.group(0).replace(",", "."))

    def _field_for_header(self, value) -> str | None:
        header = _norm(value)
        if not header:
            return None
        for field, aliases in NORMALIZED_ALIASES.items():
            if header in aliases:
                return field
        return None

    def _detect_header_row(self, ws) -> tuple[int, dict[int, str], dict[int, str]]:
        best_row = 0
        best_map: dict[int, str] = {}
        best_labels: dict[int, str] = {}
        for row_idx in range(1, min(ws.max_row, 35) + 1):
            mapping: dict[int, str] = {}
            labels: dict[int, str] = {}
            seen_fields: set[str] = set()
            for col_idx in range(1, min(ws.max_column, 60) + 1):
                raw = ws.cell(row=row_idx, column=col_idx).value
                labels[col_idx] = str(raw or "").strip()
                field = self._field_for_header(raw)
                if field and field not in seen_fields:
                    mapping[col_idx] = field
                    seen_fields.add(field)
            if len(mapping) > len(best_map):
                best_row, best_map, best_labels = row_idx, mapping, labels
        if len(best_map) < 2:
            raise ValueError("JAM could not identify enough known columns in this Excel file.")
        return best_row, best_map, best_labels

    def _payload_for_row(self, ws, row_idx: int, mapping: dict[int, str]) -> dict:
        payload: dict = {}
        for col_idx, field in mapping.items():
            raw = ws.cell(row=row_idx, column=col_idx).value
            if field == "rating":
                payload[field] = self._rating(raw)
            elif field == "match_score":
                payload[field] = self._score(raw)
            elif field in {"salary_min", "salary_max"}:
                payload[field] = self._number(raw)
            else:
                payload[field] = self._value_to_text(raw)
        payload.setdefault("status", "Saved")
        if not str(payload.get("status") or "").strip():
            payload["status"] = "Saved"
        return payload

    @staticmethod
    def _meaningful(payload: dict) -> bool:
        return any(
            str(payload.get(key) or "").strip()
            for key in ("company", "job_title", "url", "description", "source", "location")
        )

    def inspect_file(self, path: str | Path) -> dict:
        candidate = Path(path)
        if not candidate.exists():
            raise FileNotFoundError("Excel file not found.")
        workbook = load_workbook(candidate, data_only=True, read_only=True)
        ws = workbook.active
        header_row, mapping, labels = self._detect_header_row(ws)

        total_rows = 0
        duplicates = 0
        preview = []
        for row_idx in range(header_row + 1, ws.max_row + 1):
            payload = self._payload_for_row(ws, row_idx, mapping)
            if not self._meaningful(payload):
                continue
            total_rows += 1
            finder = getattr(self.repo, "find_duplicates", None)
            found = finder(payload) if callable(finder) else []
            if found:
                duplicates += 1
            if len(preview) < 8:
                preview.append({
                    "row": row_idx,
                    "company": payload.get("company", ""),
                    "job_title": payload.get("job_title", ""),
                    "status": payload.get("status", "Saved"),
                    "duplicate": bool(found),
                })

        columns = []
        for col_idx in range(1, ws.max_column + 1):
            header = labels.get(col_idx, "")
            field = mapping.get(col_idx)
            if not header and not field:
                continue
            columns.append({
                "column": col_idx,
                "header": header or f"Column {col_idx}",
                "field": field,
                "recognized": bool(field),
            })

        return {
            "path": str(candidate),
            "file_name": candidate.name,
            "sheet": ws.title,
            "header_row": header_row,
            "total_rows": total_rows,
            "duplicate_rows": duplicates,
            "mapped_fields": sorted(set(mapping.values())),
            "unknown_columns": [item["header"] for item in columns if not item["recognized"]],
            "columns": columns,
            "preview": preview,
        }

    def import_file(self, path: str | Path, import_duplicates: bool = False) -> dict:
        candidate = Path(path)
        if not candidate.exists():
            raise FileNotFoundError("Excel file not found.")
        workbook = load_workbook(candidate, data_only=True)
        ws = workbook.active
        header_row, mapping, _labels = self._detect_header_row(ws)

        imported = 0
        skipped = 0
        duplicate_skipped = 0
        failures = []

        for row_idx in range(header_row + 1, ws.max_row + 1):
            try:
                payload = self._payload_for_row(ws, row_idx, mapping)
                if not self._meaningful(payload):
                    skipped += 1
                    continue

                finder = getattr(self.repo, "find_duplicates", None)
                duplicates = finder(payload) if callable(finder) else []
                if duplicates and not import_duplicates:
                    duplicate_skipped += 1
                    continue

                self.repo.create(payload)
                imported += 1
            except Exception as exc:
                failures.append({"row": row_idx, "error": str(exc)})

        return {
            "path": str(candidate),
            "sheet": ws.title,
            "header_row": header_row,
            "mapped_fields": sorted(set(mapping.values())),
            "imported": imported,
            "skipped": skipped,
            "duplicates_skipped": duplicate_skipped,
            "failed": len(failures),
            "failures": failures,
        }
