from pathlib import Path

from openpyxl import Workbook, load_workbook

from services.export_service import ExportService
from services.import_service import ExcelImportService


class FakeRepository:
    def __init__(self):
        self.created = []

    def create(self, payload):
        self.created.append(dict(payload))
        return payload


class TempExportService(ExportService):
    def __init__(self, tmp_path: Path):
        self.tmp_path = tmp_path

    def _filename(self, suffix: str) -> Path:
        return self.tmp_path / f"export.{suffix}"

    def _track(self, export_type: str, path: Path, row_count: int) -> None:
        return None


def test_excel_import_matches_columns_out_of_order(tmp_path: Path):
    path = tmp_path / "applications.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Source", "Date", "Job Title", "Company", "Rating", "Location", "Status"])
    ws.append(["LinkedIn", "2026-09-22", "Data Analyst", "AYcode", "★★★★", "Tunis", "Applied"])
    wb.save(path)

    repo = FakeRepository()
    service = ExcelImportService(repo)
    result = service.import_file(path)

    assert result["imported"] == 1
    assert repo.created[0]["company"] == "AYcode"
    assert repo.created[0]["job_title"] == "Data Analyst"
    assert repo.created[0]["source"] == "LinkedIn"
    assert repo.created[0]["rating"] == 4
    assert repo.created[0]["status"] == "Applied"


def test_excel_export_contains_branding_filters_and_table(tmp_path: Path):
    service = TempExportService(tmp_path)
    rows = [
        {
            "date_saved": "2026-09-22T10:00:00+02:00",
            "company": "Damergi",
            "job_title": "Data Analyst",
            "status": "Applied",
            "match_score": 82,
            "rating": 5,
            "location": "Tunis",
            "work_mode": "Hybrid",
            "source": "LinkedIn",
            "date_applied": "2026-09-22T10:30:00+02:00",
            "contact_name": "",
            "contact_url": "",
            "url": "https://example.com/job",
            "notes": "Strong fit",
        }
    ]

    path = Path(
        service.export_excel(
            rows,
            filter_summary=["Status: Applied", "Location: Tunis"],
        )
    )

    wb = load_workbook(path, data_only=True)
    ws = wb["Applications"]
    assert ws["A1"].value == "JAM — Job Application Manager"
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "FILTERS APPLIED" in values
    assert "Company" in values
    assert "Damergi" in values
    assert "Data Analyst" in values
