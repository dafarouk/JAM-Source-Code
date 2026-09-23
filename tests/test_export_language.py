from pathlib import Path

from openpyxl import load_workbook

from services.export_service import ExportService


class FrenchSettings:
    def get(self, key: str, default: str = "") -> str:
        if key == "language":
            return "fr"
        return default


class TempFrenchExportService(ExportService):
    def __init__(self, tmp_path: Path):
        self.tmp_path = tmp_path
        self._settings = FrenchSettings()

    def _filename(self, suffix: str) -> Path:
        return self.tmp_path / f"export_fr.{suffix}"

    def _track(self, export_type: str, path: Path, row_count: int) -> None:
        return None


def test_french_excel_export_translates_visible_report_text(tmp_path: Path):
    service = TempFrenchExportService(tmp_path)
    rows = [
        {
            "date_saved": "2026-09-22T10:00:00+02:00",
            "company": "Damergi",
            "job_title": "Data Analyst",
            "status": "Applied",
            "match_score": 82,
            "rating": 5,
            "location": "Tunis",
            "work_mode": "Hybrid",
            "source": "LinkedIn",
            "date_applied": "2026-09-22T10:30:00+02:00",
            "contact_name": "",
            "contact_url": "",
            "url": "https://example.com/job",
            "notes": "Strong fit",
        }
    ]

    path = Path(
        service.export_excel(
            rows,
            filter_summary=["Status: Applied"],
        )
    )

    workbook = load_workbook(path, data_only=True)
    assert "Candidatures" in workbook.sheetnames
    assert "Résumé" in workbook.sheetnames

    sheet = workbook["Candidatures"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]

    assert "EXPORT DES CANDIDATURES" in values
    assert "FILTRES APPLIQUÉS" in values
    assert "Entreprise" in values
    assert "Poste" in values
    assert "Candidature envoyée" in values
